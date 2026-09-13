"""Xuất kết quả bộ thử 10.000 câu ra Excel để người dùng tự soi và chấm lại.

    .venv/bin/python scripts/xuat_excel_bo_thu.py --ket-qua <jsonl> [--truoc <jsonl>] --ra <xlsx>

Mọi câu của bộ câu (sinh lại theo --so/--seed) đều có một dòng, kể cả câu chưa
chạy. `--truoc` là kết quả của bản code cũ trên CÙNG bộ câu, để so từng câu.
"""
import argparse
import importlib.util
import re
import json
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

GOC = Path(__file__).resolve().parents[1]
_spec = importlib.util.spec_from_file_location("bo_thu", GOC / "scripts" / "bo_thu_10k_tai_lieu.py")
B = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(B)

TEN_SP = {"vay_tin_chap": "Vay tín chấp", "vay_mua_nha": "Vay mua nhà",
          "the_tin_dung": "Thẻ tín dụng", "tiet_kiem": "Gửi tiết kiệm", "": "Chung"}
TEN_LOAI = {"fact": "Hỏi thông tin", "che": "Khách chê", "danh_muc": "Hỏi danh mục",
            "ngoai": "Ngoài tài liệu", "tinh": "Tính trả góp", "nhu_cau": "Nêu nhu cầu vay"}
TEN_NGUON = {"llm": "Mô hình AI", "tra_tu_quy_tac_tai_chinh": "Luật tài chính",
             "tra_tu_danh_muc": "Danh mục sản phẩm", "tra_tu_ho_so": "Hồ sơ khách",
             "luot_thuong_gap": "Lượt thường gặp", "bang_doc_thang": "Bảng hỏi-đáp",
             "llm_rong": "Mô hình trả rỗng"}
TEN_LOI = {"rong": "Trả lời rỗng", "chu_ngoai": "Lọt chữ nước ngoài",
           "thieu_dap_an": "Thiếu thông tin đúng", "hua_kiem_tra": "Hứa kiểm tra thay vì trả lời",
           "lac_san_pham": "Đọc số của sản phẩm khác", "so_ngoai_tai_lieu": "Số không có trong tài liệu",
           "bia_ngoai_pham_vi": "Bịa số cho câu ngoài tài liệu",
           "sai_huong": "Sai hướng hạn mức", "qua_dai": "Dài quá 60 từ (chỉ cảnh báo)",
           "loi_ket_noi": "Lỗi kết nối"}

MUC = "1F4E5F"      # xanh cổ vịt đậm: tiêu đề
NHAT = "E8EFF1"     # nền nhãn
DAT = "DCEFE3"
TRUOT = "F7DCDA"
CHUA = "EFEFEF"
VIEN = Border(bottom=Side(style="thin", color="C9D3D6"))


def doc(tep):
    if not tep or not Path(tep).exists():
        return {}
    ra = {}
    for x in Path(tep).read_text(encoding="utf-8").splitlines():
        if x.strip():
            d = json.loads(x)
            ra[d["i"]] = d
    return ra


def hong(d):
    return any(e in B.LOI_TRUOT or e == "loi_ket_noi" for e in d["loi"])


def ket_luan(d):
    if d is None:
        return "Chưa chạy"
    return "Trượt" if hong(d) else "Đạt"


def dap_an_de_doc(q):
    if q["loai"] == "tinh":
        if q["x"] > B.TRAN_TC:
            return "Nói rõ vượt hạn mức 500 triệu"
        ky = q["x"] / q["t"] + q["x"] * B.LAI_TC / 100 / 12
        return f"Tháng đầu khoảng {ky:.1f} triệu (±10%)".replace(".", ",")
    if q["loai"] == "nhu_cau":
        return "Nói vượt hạn mức" if q["x"] > B.TRAN_TC else "Không được nói vượt hạn mức"
    if q["loai"] == "ngoai":
        return "Không bịa con số"
    cac = []
    for r in q.get("dap") or []:
        s = (r.replace("\\b", "").replace(" ?", " ").replace("[.,]", ",").replace("(?:", "(")
             .replace(".{0,15}", "…").replace(".{0,12}", "…").replace(".{0,4}", "…"))
        s = re.sub(r"\.\{0,\d+\}", "…", s.replace(",?", ","))
        s = s.replace("\\d", "x").replace("\\s", " ").replace("?", "")
        s = s.replace("(", "").replace(")", "").replace("|", " hoặc ")
        s = re.sub(r"KT 3", "KT3", s)
        cac.append(s)
    if len(cac) > 1 and q["loai"] == "danh_muc" and q["chu_de"].startswith("danh_muc"):
        return "Đủ cả: " + "; ".join(cac)
    return "Có một trong: " + "; ".join(cac) if len(cac) > 1 else (cac[0] if cac else "")


def tieu_de(ws, hang, cot_ten):
    for c, ten in enumerate(cot_ten, 1):
        o = ws.cell(hang, c, ten)
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = PatternFill("solid", fgColor=MUC)
        o.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[hang].height = 30


def bang_nho(ws, hang, ten_bang, cot_ten, dong, dinh_dang=None):
    ws.cell(hang, 1, ten_bang).font = Font(bold=True, size=12, color=MUC)
    hang += 1
    tieu_de(ws, hang, cot_ten)
    for r in dong:
        hang += 1
        for c, v in enumerate(r, 1):
            o = ws.cell(hang, c, v)
            o.border = VIEN
            if dinh_dang and c in dinh_dang:
                o.number_format = dinh_dang[c]
    return hang + 2


def thong_ke(ket_qua, cau, khoa):
    b = defaultdict(lambda: [0, 0])
    for i, d in ket_qua.items():
        k = khoa(cau[i], d)
        b[k][0] += 1
        b[k][1] += hong(d)
    return b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ket-qua", required=True)
    ap.add_argument("--truoc", default="")
    ap.add_argument("--ra", required=True)
    ap.add_argument("--so", type=int, default=10000)
    ap.add_argument("--seed", type=int, default=13)
    ap.add_argument("--nhan", default="Bản đã sửa luật theo sản phẩm")
    a = ap.parse_args()

    cau = {q["i"]: q for q in B.sinh_cau_hoi(a.so, a.seed)}
    sau, truoc = doc(a.ket_qua), doc(a.truoc)

    wb = Workbook()
    tq = wb.active
    tq.title = "Tổng quan"
    ct = wb.create_sheet("Tất cả câu hỏi")
    cd = wb.create_sheet("Theo chủ đề")
    tr = wb.create_sheet("Câu trượt")

    # --- Tất cả câu hỏi ------------------------------------------------------
    cot = ["STT", "Sản phẩm", "Chủ đề", "Loại câu", "Câu khách hỏi", "Lỗi nghe nhầm cài vào",
           "Kiểu gửi", "Câu AI trả lời", "Nguồn trả lời", "Máy chấm", "Lỗi máy bắt",
           "Thông tin phải có", "Thời gian (ms)"]
    if truoc:
        cot += ["Bản cũ: máy chấm", "Bản cũ: câu trả lời"]
    cot += ["Bạn chấm", "Ghi chú của bạn"]
    tieu_de(ct, 1, cot)
    tieu_de(tr, 1, cot)
    rong = {"STT": 7, "Sản phẩm": 14, "Chủ đề": 18, "Loại câu": 15, "Câu khách hỏi": 42,
            "Lỗi nghe nhầm cài vào": 16, "Kiểu gửi": 10, "Câu AI trả lời": 60, "Nguồn trả lời": 17,
            "Máy chấm": 11, "Lỗi máy bắt": 28, "Thông tin phải có": 30, "Thời gian (ms)": 11,
            "Bản cũ: máy chấm": 12, "Bản cũ: câu trả lời": 50, "Bạn chấm": 12, "Ghi chú của bạn": 30}
    for ws in (ct, tr):
        for c, ten in enumerate(cot, 1):
            ws.column_dimensions[get_column_letter(c)].width = rong[ten]
        ws.freeze_panes = "F2"

    hang_tr = 1
    for i in sorted(cau):
        q, d, d0 = cau[i], sau.get(i), truoc.get(i)
        dong = [i + 1, TEN_SP.get(q["sp"], q["sp"]), q["chu_de"], TEN_LOAI.get(q["loai"], q["loai"]),
                (d or {}).get("cau") or q["cau"], (d or q).get("nghe_nham") or "",
                (d or {}).get("kieu", "").replace("text_soi", "chữ").replace("text", "chữ + câu đệm"),
                (d or {}).get("tra_loi", ""),
                TEN_NGUON.get(((d or {}).get("nguon") or "").split(":")[0], (d or {}).get("nguon", "")),
                ket_luan(d), "; ".join(TEN_LOI.get(e, e) for e in (d or {}).get("loi", [])),
                dap_an_de_doc(q), (d or {}).get("ms") or None]
        if truoc:
            dong += [ket_luan(d0), (d0 or {}).get("tra_loi", "")]
        dong += [None, None]
        ct.append(dong)
        if d and hong(d):
            tr.append(dong)
            hang_tr += 1

    n_cot = len(cot)
    cot_cham = get_column_letter(cot.index("Máy chấm") + 1)
    for ws, cuoi in ((ct, len(cau) + 1), (tr, hang_tr)):
        if cuoi < 2:
            continue
        vung = f"A2:{get_column_letter(n_cot)}{cuoi}"
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cot)}{cuoi}"
        for mau, chu in ((DAT, "Đạt"), (TRUOT, "Trượt"), (CHUA, "Chưa chạy")):
            ws.conditional_formatting.add(vung, FormulaRule(
                formula=[f'${cot_cham}2="{chu}"'], fill=PatternFill("solid", fgColor=mau)))
        dv = DataValidation(type="list", formula1='"Đúng,Sai,Chưa chắc"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(cot.index('Bạn chấm') + 1)}2:"
               f"{get_column_letter(cot.index('Bạn chấm') + 1)}{cuoi}")
        for hang in ws.iter_rows(min_row=2, max_row=cuoi):
            for o in hang:
                o.alignment = Alignment(vertical="top", wrap_text=True)

    # --- Theo chủ đề ----------------------------------------------------------
    b_sau = thong_ke(sau, cau, lambda q, d: (q["sp"], q["chu_de"], q["loai"]))
    b_truoc = thong_ke(truoc, cau, lambda q, d: (q["sp"], q["chu_de"], q["loai"])) if truoc else {}
    cot_cd = ["Sản phẩm", "Chủ đề", "Loại câu", "Số câu", "Đạt", "Trượt", "Tỉ lệ đạt"]
    if truoc:
        cot_cd += ["Bản cũ: số câu", "Bản cũ: tỉ lệ đạt"]
    tieu_de(cd, 1, cot_cd)
    for k, (so, t) in sorted(b_sau.items(), key=lambda kv: kv[1][1] / kv[1][0], reverse=True):
        dong = [TEN_SP.get(k[0], k[0]), k[1], TEN_LOAI.get(k[2], k[2]), so, so - t, t, (so - t) / so]
        if truoc:
            s0, t0 = b_truoc.get(k, [0, 0])
            dong += [s0 or None, (s0 - t0) / s0 if s0 else None]
        cd.append(dong)
    for c, w in enumerate([15, 22, 16, 9, 9, 9, 11, 13, 14], 1):
        cd.column_dimensions[get_column_letter(c)].width = w
    for hang in cd.iter_rows(min_row=2):
        hang[6].number_format = "0.0%"
        if truoc:
            hang[8].number_format = "0.0%"
    cd.freeze_panes = "A2"
    cd.auto_filter.ref = f"A1:{get_column_letter(len(cot_cd))}{cd.max_row}"
    cd.conditional_formatting.add(f"G2:G{cd.max_row}", FormulaRule(
        formula=["G2<0.8"], fill=PatternFill("solid", fgColor=TRUOT)))

    # --- Tổng quan ------------------------------------------------------------
    tq.column_dimensions["A"].width = 34
    for c in "BCDEF":
        tq.column_dimensions[c].width = 16
    tq["A1"] = "Bộ thử 10.000 câu hỏi về tài liệu"
    tq["A1"].font = Font(bold=True, size=16, color=MUC)
    tq["A2"] = (f"{a.nhan} · seed {a.seed} · xuất lúc {datetime.now():%H:%M %d-%m-%Y} · "
                "mỗi câu một phiên mới, máy chủ thật trên máy Windows")
    tq["A2"].font = Font(italic=True, color="5A6B70")

    so_chay = len(sau)
    so_truot = sum(hong(d) for d in sau.values())
    ms = [d["ms"] for d in sau.values() if d.get("ms")]
    hang = 4
    kpi = [("Tổng câu trong bộ", len(cau), None), ("Đã chạy", so_chay, None),
           ("Máy chấm đạt", so_chay - so_truot, None), ("Máy chấm trượt", so_truot, None),
           ("Tỉ lệ đạt", (so_chay - so_truot) / so_chay if so_chay else None, "0.0%"),
           ("Thời gian trả lời trung vị (ms)", statistics.median(ms) if ms else None, None),
           ("Thời gian trả lời p95 (ms)", sorted(ms)[int(len(ms) * .95)] if ms else None, None)]
    cot_ban = get_column_letter(cot.index("Bạn chấm") + 1)
    kpi += [("Bạn đã chấm", f"=COUNTA('Tất cả câu hỏi'!{cot_ban}2:{cot_ban}{len(cau) + 1})", None),
            ("Bạn chấm Sai", f"=COUNTIF('Tất cả câu hỏi'!{cot_ban}2:{cot_ban}{len(cau) + 1},\"Sai\")", None)]
    for ten, v, fmt in kpi:
        tq.cell(hang, 1, ten).fill = PatternFill("solid", fgColor=NHAT)
        o = tq.cell(hang, 2, v)
        o.font = Font(bold=True)
        if fmt:
            o.number_format = fmt
        hang += 1
    hang += 1

    def bang_tl(ten, khoa, tu_dien):
        b = thong_ke(sau, cau, khoa)
        b0 = thong_ke(truoc, cau, khoa) if truoc else {}
        cot_b = ["", "Số câu", "Trượt", "Tỉ lệ đạt"] + (["Bản cũ: số câu", "Bản cũ: tỉ lệ đạt"] if truoc else [])
        dong = []
        for k, (so, t) in sorted(b.items(), key=lambda kv: -kv[1][0]):
            r = [tu_dien.get(k, k) if tu_dien else k, so, t, (so - t) / so]
            if truoc:
                s0, t0 = b0.get(k, [0, 0])
                r += [s0 or None, (s0 - t0) / s0 if s0 else None]
            dong.append(r)
        return bang_nho(tq, hang, ten, cot_b, dong, {4: "0.0%", 6: "0.0%"})

    hang = bang_tl("Theo sản phẩm", lambda q, d: q["sp"], TEN_SP)
    hang = bang_tl("Theo nguồn trả lời", lambda q, d: d["nguon"].split(":")[0], TEN_NGUON)
    hang = bang_tl("Theo loại câu", lambda q, d: q["loai"], TEN_LOAI)
    hang = bang_tl("Có lỗi nghe nhầm cài vào", lambda q, d: "Có" if d.get("nghe_nham") else "Không", None)
    dem_loi = Counter(e for d in sau.values() for e in d["loi"])
    hang = bang_nho(tq, hang, "Lỗi máy bắt (một câu có thể dính nhiều lỗi)", ["Lỗi", "Số câu"],
                    [[TEN_LOI.get(e, e), n] for e, n in dem_loi.most_common()])
    tq.cell(hang, 1, "Cách dùng: lọc cột “Máy chấm” ở trang Tất cả câu hỏi, đọc câu trả lời, "
                     "chọn Đúng/Sai ở cột “Bạn chấm”. Trang Tổng quan tự đếm số câu bạn đã chấm.")
    tq.cell(hang, 1).font = Font(italic=True, color="5A6B70")

    wb.save(a.ra)
    print(f"đã xuất {a.ra}: {len(cau)} câu, đã chạy {so_chay}, trượt {so_truot}")


if __name__ == "__main__":
    main()
