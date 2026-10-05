#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import statistics
from pathlib import Path

from openpyxl import load_workbook

from common import has_cjk_hangul, vietnamese_score


def main() -> None:
    ap = argparse.ArgumentParser(description="Đổi Excel benchmark 10k thành metrics cho BankVN gate")
    ap.add_argument("excel")
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    wb = load_workbook(args.excel, read_only=True, data_only=True)
    if "Tất cả câu hỏi" not in wb.sheetnames:
        raise SystemExit("[ERROR] Excel không có sheet 'Tất cả câu hỏi'")
    ws = wb["Tất cả câu hỏi"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(v or "").strip() for v in next(rows)]
    idx = {name: i for i, name in enumerate(headers)}
    for needed in ("Máy chấm", "Câu AI trả lời"):
        if needed not in idx:
            raise SystemExit(f"[ERROR] Thiếu cột {needed!r}")

    passed = failed = answered = vi_ok = foreign = 0
    times: list[float] = []
    for row in rows:
        status = str(row[idx["Máy chấm"]] or "").strip()
        if status == "Đạt":
            passed += 1
        elif status == "Trượt":
            failed += 1
        answer = str(row[idx["Câu AI trả lời"]] or "").strip()
        if answer:
            answered += 1
            vi_ok += vietnamese_score(answer) >= 0.05
            foreign += has_cjk_hangul(answer)
        if "Thời gian (ms)" in idx:
            value = row[idx["Thời gian (ms)"]]
            if isinstance(value, (int, float)) and value >= 0:
                times.append(float(value))

    total = passed + failed
    metrics = {
        "domain_accuracy": passed / total if total else 0.0,
        "vietnamese_rate": vi_ok / answered if answered else 0.0,
        "foreign_script_rate": foreign / answered if answered else 0.0,
        "questions_scored": total,
        "answered": answered,
    }
    if times:
        ordered = sorted(times)
        pos = min(len(ordered) - 1, max(0, round(0.95 * (len(ordered) - 1))))
        metrics["p95_latency_ms"] = ordered[pos]
        metrics["median_latency_ms"] = statistics.median(ordered)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(metrics, ensure_ascii=False))


if __name__ == "__main__":
    main()
